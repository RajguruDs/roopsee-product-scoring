"""Validation for the V3 ingredient score master. Data hygiene only - no scoring logic.

WHY THIS EXISTS
---------------
`read_ingredient_scores()` (tools/build_automated_scores.py:791) has two silent failure
modes that the master audit surfaced:

  1. A BLANK cell becomes 40 (`:816`). 40 is not "unknown" - the workbook's own `Scale`
     sheet defines it as "Tolerated, no particular benefit", a real clinical statement.
     A missing rating and a deliberate "tolerated" rating therefore become
     indistinguishable, and a row the author simply had not filled in would score as if
     they had. The master has ZERO blanks today, so this is latent, not active - but it
     is silent, which is the part worth fixing.

  2. `score_value()` (`:567`) only rounds. A typo such as 900, or a stray 75, would flow
     straight through into scoring with no complaint.

This module validates rather than repairs. It never rewrites a score, because a score is
a clinical statement that only the master's author can make.

WHAT IT DOES NOT DO
-------------------
It does not change the V3 formula, aggregation, the 90 ceiling, Primary/Secondary tier
logic, ranking, the UI, the production scorer or the canonical product population. It is
imported by the V3 experimental scorer only; `build_automated_scores.py` and the
production build path are untouched.
"""

from __future__ import annotations

from pathlib import Path
from typing import Any

# The value contract, taken verbatim from the workbook's own "Scale" sheet.
#   -100 Disqualifying - should gate the product out, not average in
#      0 Unsuitable / no benefit
#     40 Tolerated, no particular benefit
#     50 Neutral (age and pregnancy/breastfeeding columns)
#     90 Well suited
#    100 Ideal / fully cleared
MASTER_VALUE_CONTRACT: frozenset[float] = frozenset({-100.0, 0.0, 40.0, 50.0, 90.0, 100.0})

# Derived columns computed by the loader (`recompute_unknown_skin_scores`) rather than
# authored, so they legitimately hold values outside the contract and are not validated.
DERIVED_COLUMNS: frozenset[str] = frozenset({"I dont know Score", "I dont know+Sensitive Score"})

# Placeholder text that reaches the ingredient stream from upstream spreadsheets. These
# are NOT ingredients: they are a human writing "there isn't one here". Left in place they
# never resolve against the master (so they contribute nothing to any score) but they do
# inflate the tier lengths that `evidence_strength` is derived from, which makes a product
# look like it claims an active it does not have.
PLACEHOLDER_LABELS: frozenset[str] = frozenset({
    "not used", "notused", "not applicable", "n/a", "na", "nil", "none", "-", "--", "",
})


class MasterValidationError(RuntimeError):
    """Raised when the ingredient master violates its own declared value contract."""


def is_placeholder(label: str) -> bool:
    """True when a label is filler text rather than an ingredient name."""
    return (label or "").strip().lower().strip(".") in PLACEHOLDER_LABELS


def strip_placeholders(tiers: dict[str, list[str]]) -> tuple[dict[str, list[str]], list[str]]:
    """Drop placeholder labels from Primary/Secondary/Incidental tiers.

    Returns the cleaned tiers plus the labels removed, so the caller can report them.
    Tier membership is otherwise untouched - this removes non-ingredients, it does not
    reclassify ingredients between tiers.
    """
    removed: list[str] = []
    cleaned: dict[str, list[str]] = {}
    for tier, names in tiers.items():
        keep = []
        for name in names or []:
            if is_placeholder(name):
                removed.append(f"{tier}:{name}")
            else:
                keep.append(name)
        cleaned[tier] = keep
    return cleaned, removed


def validate_loaded_master(master: dict[str, dict[str, Any]]) -> dict[str, Any]:
    """Check every authored score cell against the contract.

    Returns a report. Raises nothing - the caller decides whether a violation is fatal,
    because a build may want to fail while an analysis run may want to continue and list
    every problem at once.
    """
    violations: list[dict[str, Any]] = []
    for name, row in master.items():
        for column, value in (row.get("scores") or {}).items():
            if column in DERIVED_COLUMNS or column == "None":
                continue
            if value is None:
                violations.append({"ingredient": name, "column": column,
                                   "value": None, "problem": "missing value"})
            elif float(value) not in MASTER_VALUE_CONTRACT:
                violations.append({"ingredient": name, "column": column,
                                   "value": value, "problem": "outside the value contract"})
    return {
        "ingredients": len(master),
        "violations": violations,
        "ok": not violations,
        "contract": sorted(MASTER_VALUE_CONTRACT),
    }


def audit_raw_workbook(path: str | Path, header_map: dict[str, str]) -> dict[str, Any]:
    """Inspect the raw cells BEFORE the loader's `else 40` masks anything.

    This is the only place a blank can still be seen as a blank, so it is the only place
    the silent-default problem can actually be detected.
    """
    import openpyxl  # local import: only needed when auditing the source workbook

    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    sheet = workbook["Ingredient Scores"]
    rows = sheet.iter_rows(values_only=True)
    headers = [str(h).strip() if h is not None else "" for h in next(rows)]
    index = {h: i for i, h in enumerate(headers)}

    blanks: list[dict[str, Any]] = []
    off_contract: list[dict[str, Any]] = []
    non_numeric: list[dict[str, Any]] = []
    seen: dict[str, int] = {}
    duplicate_names: list[dict[str, Any]] = []
    total = 0

    for row_number, raw in enumerate(rows, start=2):
        name = str(raw[0]).strip() if raw[0] is not None else ""
        if not name:
            continue
        if name in seen:
            duplicate_names.append({"ingredient": name, "row": row_number,
                                    "first_seen_row": seen[name]})
        else:
            seen[name] = row_number
        for source_header in header_map:
            if source_header not in index:
                continue
            total += 1
            value = raw[index[source_header]]
            if value is None or (isinstance(value, str) and not value.strip()):
                blanks.append({"ingredient": name, "column": source_header, "row": row_number})
                continue
            try:
                number = float(value)
            except (TypeError, ValueError):
                non_numeric.append({"ingredient": name, "column": source_header,
                                    "row": row_number, "value": str(value)})
                continue
            if number not in MASTER_VALUE_CONTRACT:
                off_contract.append({"ingredient": name, "column": source_header,
                                     "row": row_number, "value": number})

    return {
        "workbook": str(path),
        "rows": len(seen),
        "cells_checked": total,
        # Every one of these would have become a silent 40 in the current loader.
        "blank_cells": blanks,
        "off_contract_cells": off_contract,
        "non_numeric_cells": non_numeric,
        "duplicate_canonical_names": duplicate_names,
        "ok": not (blanks or off_contract or non_numeric or duplicate_names),
    }
