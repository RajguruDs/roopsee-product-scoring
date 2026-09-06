"""Run the scoring-quality metrics and write the reports (Steps 5-10).

Reads scoring_quality_predictions.csv. Doctor scores may come either from that
file directly or from a completed doctor workbook merged in with --doctor-file.

    python tools/run_scoring_validation.py
    python tools/run_scoring_validation.py --doctor-file outputs/roopsee_canonical/doctor_validation_completed.xlsx

Writes into outputs/roopsee_canonical/:

  scoring_quality_validation_report.md      the old vs new report
  scoring_quality_metrics.json              every metric, machine-readable
  performance_by_product_type.csv           and by concern / skin type /
  performance_by_*.csv                        confidence / special condition / ...
  scoring_quality_ranking.csv               per-profile ranking agreement
  scoring_quality_failures.csv              the 50 largest canonical-vs-doctor errors

This module measures only. It never writes to data/, never changes a score, and
never invents a doctor answer: with no doctor scores present it says so and
stops rather than reporting a result.
"""

from __future__ import annotations

import argparse
import csv
import json
import os
import sys
from collections import defaultdict
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

sys.path.insert(0, str(Path(__file__).resolve().parent))
import scoring_validation as sv  # noqa: E402


REPO_ROOT = Path(__file__).resolve().parents[1]
CANONICAL_OUTPUT_DIR = Path(
    os.environ.get("ROOPSEE_CANONICAL_OUTPUT_DIR", REPO_ROOT / "outputs" / "roopsee_canonical")
)

PREDICTIONS_PATH = CANONICAL_OUTPUT_DIR / "scoring_quality_predictions.csv"
META_PATH = CANONICAL_OUTPUT_DIR / "scoring_quality_validation_meta.json"
REPORT_PATH = CANONICAL_OUTPUT_DIR / "scoring_quality_validation_report.md"
METRICS_PATH = CANONICAL_OUTPUT_DIR / "scoring_quality_metrics.json"
RANKING_PATH = CANONICAL_OUTPUT_DIR / "scoring_quality_ranking.csv"
FAILURES_PATH = CANONICAL_OUTPUT_DIR / "scoring_quality_failures.csv"

BREAKDOWNS = [
    ("product_type", "performance_by_product_type.csv"),
    ("concern", "performance_by_concern.csv"),
    ("skin_type", "performance_by_skin_type.csv"),
    ("confidence", "performance_by_confidence.csv"),
    ("mapping_confidence", "performance_by_mapping_confidence.csv"),
    ("special_conditions", "performance_by_special_condition.csv"),
    ("canonical_v2_bin", "performance_by_score_bucket.csv"),
    ("canonical_v2_hard_blocked", "performance_by_hard_block_status.csv"),
    ("profile_id", "performance_by_profile.csv"),
]


def number(value: Any) -> float | None:
    if value in (None, ""):
        return None
    try:
        return float(value)
    except (TypeError, ValueError):
        return None


def load_predictions() -> list[dict[str, Any]]:
    if not PREDICTIONS_PATH.exists():
        raise SystemExit(f"Missing {PREDICTIONS_PATH}.\nRun: python tools/build_scoring_validation.py")
    with PREDICTIONS_PATH.open(newline="", encoding="utf-8-sig") as handle:
        return list(csv.DictReader(handle))


def merge_doctor_file(rows: list[dict[str, Any]], path: Path) -> int:
    """Merge doctor answers keyed on (product, profile). Never overwrites with blanks."""
    if not path.exists():
        raise SystemExit(f"Doctor file not found: {path}")

    answers: dict[tuple[str, str], dict[str, str]] = {}
    if path.suffix.lower() in {".xlsx", ".xlsm"}:
        try:
            from openpyxl import load_workbook
        except ImportError as error:
            raise SystemExit("openpyxl is required to read an xlsx doctor file") from error
        workbook = load_workbook(path, data_only=True)
        sheet = workbook["Validation"] if "Validation" in workbook.sheetnames else workbook.active
        header = [str(cell.value or "").strip() for cell in next(sheet.iter_rows(min_row=1, max_row=1))]
        for raw in sheet.iter_rows(min_row=2, values_only=True):
            record = dict(zip(header, raw))
            key = (str(record.get("Product ID") or ""), str(record.get("Profile ID") or ""))
            if key[0] and key[1]:
                answers[key] = {
                    "doctor_score": record.get("DOCTOR SCORE"),
                    "doctor_recommendation": record.get("DOCTOR RECOMMENDATION"),
                    "doctor_notes": record.get("DOCTOR NOTES"),
                }
    else:
        with path.open(newline="", encoding="utf-8-sig") as handle:
            for record in csv.DictReader(handle):
                key = (
                    record.get("canonical_product_id_v2") or record.get("Product ID") or "",
                    record.get("profile_id") or record.get("Profile ID") or "",
                )
                if key[0] and key[1]:
                    answers[key] = {
                        "doctor_score": record.get("doctor_score") or record.get("DOCTOR SCORE"),
                        "doctor_recommendation": record.get("doctor_recommendation")
                        or record.get("DOCTOR RECOMMENDATION"),
                        "doctor_notes": record.get("doctor_notes") or record.get("DOCTOR NOTES"),
                    }

    merged = 0
    for row in rows:
        answer = answers.get((row["canonical_product_id_v2"], row["profile_id"]))
        if not answer:
            continue
        if answer.get("doctor_score") not in (None, ""):
            row["doctor_score"] = str(answer["doctor_score"])
            row["doctor_recommendation"] = str(answer.get("doctor_recommendation") or "")
            row["doctor_notes"] = str(answer.get("doctor_notes") or "")
            merged += 1
    return merged


def write_csv(path: Path, rows: list[dict[str, Any]]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    if not rows:
        path.write_text("", encoding="utf-8")
        return
    with path.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.DictWriter(handle, fieldnames=list(rows[0].keys()))
        writer.writeheader()
        writer.writerows(rows)


# --------------------------------------------------------------------------
# Step 7 -- ranking
# --------------------------------------------------------------------------


def ranking_report(rows: list[dict[str, Any]]) -> list[dict[str, Any]]:
    """Per-profile ranking agreement between each system and the doctor.

    Ranks are computed within the validation sample, not the full catalogue.
    A sample of ~150 products is a different shelf from 4,037, so these numbers
    describe relative ordering, not the live top-of-page experience.
    """
    by_profile: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for row in rows:
        by_profile[row["profile_id"]].append(row)

    output: list[dict[str, Any]] = []
    for profile_id in sorted(by_profile):
        group = by_profile[profile_id]
        doctor_scores = {
            row["canonical_product_id_v2"]: number(row["doctor_score"])
            for row in group
            if number(row["doctor_score"]) is not None
        }
        if not doctor_scores:
            continue

        rated = [row for row in group if row["canonical_product_id_v2"] in doctor_scores]
        doctor_ranked = [
            row["canonical_product_id_v2"]
            for row in sorted(
                rated,
                key=lambda r: (-(number(r["doctor_score"]) or 0), r["canonical_product_id_v2"]),
            )
        ]

        entry: dict[str, Any] = {
            "profile_id": profile_id,
            "profile_label": group[0]["profile_label"],
            "rated_products": len(doctor_scores),
        }

        for system, field in [("canonical", "canonical_v2_score"), ("legacy", "legacy_score")]:
            scored = [row for row in rated if number(row[field]) is not None]
            system_ranked = [
                row["canonical_product_id_v2"]
                for row in sorted(
                    scored, key=lambda r: (-(number(r[field]) or 0), r["canonical_product_id_v2"])
                )
            ]
            doctor_subset = [uid for uid in doctor_ranked if uid in set(system_ranked)]
            entry[f"{system}_scored_products"] = len(system_ranked)
            entry[f"{system}_top1_agreement"] = (
                sv.top_1_agreement(system_ranked, doctor_subset) if doctor_subset else None
            )
            entry[f"{system}_top5_overlap_pct"] = sv.top_n_overlap(system_ranked, doctor_subset, 5)
            entry[f"{system}_top10_overlap_pct"] = sv.top_n_overlap(system_ranked, doctor_subset, 10)
            entry[f"{system}_rank_spearman"] = sv.spearman_correlation(
                [number(row[field]) for row in scored],
                [number(row["doctor_score"]) for row in scored],
            )
            entry[f"{system}_unsuitable_in_top5"] = sv.unsuitable_in_top_n(system_ranked, doctor_scores, 5)
            entry[f"{system}_missed_strong_in_top10"] = len(
                sv.missed_strong_recommendations(system_ranked, doctor_scores, 10)
            )
        output.append(entry)
    return output


# --------------------------------------------------------------------------
# Step 8 -- failure analysis
# --------------------------------------------------------------------------


def failure_analysis(rows: list[dict[str, Any]], limit: int = 50) -> list[dict[str, Any]]:
    scored = []
    for row in rows:
        doctor = number(row["doctor_score"])
        canonical = number(row["canonical_v2_score"])
        if doctor is None or canonical is None:
            continue
        legacy = number(row["legacy_score"])
        category, note = sv.classify_failure(row)
        scored.append(
            {
                "canonical_product_id_v2": row["canonical_product_id_v2"],
                "gtin": row["gtin"],
                "product_name": row["product_name"],
                "brand": row["brand"],
                "product_type": row["product_type"],
                "profile_id": row["profile_id"],
                "profile_label": row["profile_label"],
                "doctor_score": doctor,
                "canonical_v2_score": canonical,
                "absolute_error": round(abs(doctor - canonical), 2),
                "legacy_score": legacy if legacy is not None else "",
                "legacy_absolute_error": round(abs(doctor - legacy), 2) if legacy is not None else "",
                "confidence": row["confidence"],
                "mapping_confidence": row["mapping_confidence"],
                "primary_ingredients": row["primary_ingredients"],
                "secondary_ingredients": row["secondary_ingredients"],
                "doctor_anchor_support": row["anchor_support"],
                "exact_support": row["exact_support"],
                "layer_baseline": row["layer_baseline"],
                "layer_v2": row["layer_v2"],
                "layer_anchor": row["layer_anchor"],
                "layer_type_family": row["layer_type_family"],
                "layer_type": row["layer_type"],
                "hard_block": row["canonical_v2_hard_blocked"],
                "doctor_hard_block": "true" if doctor <= sv.HARD_BLOCK else "false",
                "suspected_failure_category": category,
                "notes": note,
                "doctor_notes": row.get("doctor_notes", ""),
            }
        )
    scored.sort(key=lambda entry: (-entry["absolute_error"], entry["canonical_product_id_v2"]))
    return scored[:limit]


# --------------------------------------------------------------------------
# Step 9 -- known failure modes
# --------------------------------------------------------------------------


def known_failure_modes(rows: list[dict[str, Any]]) -> list[dict[str, Any]]:
    """Explicit checks on the historical problem areas, per system."""

    def collect(predicate, field: str) -> list[dict[str, Any]]:
        output = []
        for row in rows:
            doctor = number(row["doctor_score"])
            system = number(row[field])
            if doctor is None or system is None:
                continue
            if predicate(row, doctor, system):
                output.append(row)
        return output

    pregnancy = lambda row: "Pregnant" in row["special_conditions"]  # noqa: E731
    breastfeeding = lambda row: "Breastfeeding" in row["special_conditions"]  # noqa: E731
    teen = lambda row: row["age"] == "Teen"  # noqa: E731
    dryness = lambda row: "Excessive Dryness" in row["special_conditions"]  # noqa: E731

    checks = [
        ("pregnancy false-safe", lambda r, d, s: pregnancy(r) and d <= sv.HARD_BLOCK < s),
        ("pregnancy false-block", lambda r, d, s: pregnancy(r) and s <= sv.HARD_BLOCK < d),
        ("breastfeeding false-safe", lambda r, d, s: breastfeeding(r) and d <= sv.HARD_BLOCK < s),
        ("breastfeeding false-block", lambda r, d, s: breastfeeding(r) and s <= sv.HARD_BLOCK < d),
        ("teen false-safe", lambda r, d, s: teen(r) and d <= sv.HARD_BLOCK < s),
        ("teen under-scoring", lambda r, d, s: teen(r) and s > sv.HARD_BLOCK and d - s >= 20),
        ("excessive dryness false-block", lambda r, d, s: dryness(r) and s <= sv.HARD_BLOCK < d),
        ("excessive dryness under-scoring", lambda r, d, s: dryness(r) and s > sv.HARD_BLOCK and d - s >= 20),
        ("hard-block propagation disagreement", lambda r, d, s: (d <= sv.HARD_BLOCK) != (s <= sv.HARD_BLOCK)),
        (
            "low-confidence artificial cap",
            lambda r, d, s: r["confidence"] in {"Low", "Medium"} and d - s >= 15,
        ),
        ("doctor strong, system weak", lambda r, d, s: d >= 80 and s < 60),
        ("doctor weak, system strong", lambda r, d, s: d < 50 and s >= 80),
    ]

    results = []
    for label, predicate in checks:
        canonical_hits = collect(predicate, "canonical_v2_score")
        legacy_hits = collect(predicate, "legacy_score")
        results.append(
            {
                "failure_mode": label,
                "canonical_v2_cases": len(canonical_hits),
                "legacy_cases": len(legacy_hits),
                "canonical_examples": "; ".join(
                    f"{row['product_name'][:40]} [{row['profile_id']}]" for row in canonical_hits[:3]
                ),
            }
        )
    return results


# --------------------------------------------------------------------------
# Step 10 -- report
# --------------------------------------------------------------------------


def format_value(value: Any) -> str:
    if value is None:
        return "—"
    if isinstance(value, bool):
        return "yes" if value else "no"
    if isinstance(value, float):
        return f"{value:g}"
    return str(value)


def build_report(
    meta: dict[str, Any],
    comparison: list[dict[str, Any]],
    legacy: dict[str, Any],
    canonical: dict[str, Any],
    breakdown_files: dict[str, list[dict[str, Any]]],
    ranking: list[dict[str, Any]],
    failures: list[dict[str, Any]],
    modes: list[dict[str, Any]],
    doctor_count: int,
    total_rows: int,
    comparable_rows: int,
) -> str:
    lines: list[str] = [
        "# Scoring Quality Validation Report",
        "",
        f"- Validation version: `{meta.get('scoring_validation_version')}`",
        f"- Sample version: `{meta.get('validation_sample_version')}`  |  Profile version: `{meta.get('validation_profile_version')}`",
        f"- Generated: {datetime.now(timezone.utc).isoformat()}",
        f"- Scoring engine: {meta.get('scoring_engine')}",
        f"- Sample size: {meta.get('sample_size')} products x {meta.get('profile_count')} profiles = {total_rows} pairs",
        f"- Doctor scores present: **{doctor_count}** of {total_rows} pairs",
        f"- Pairs comparable across BOTH systems and the doctor: **{comparable_rows}**",
        "",
    ]

    if doctor_count == 0:
        lines += [
            "## Status: awaiting doctor review",
            "",
            "No doctor scores are present, so no metric can be computed and no conclusion",
            "about either system is possible. This is the expected state until",
            "`doctor_validation_template.xlsx` has been filled in and merged.",
            "",
            "Nothing below has been estimated, inferred or filled in automatically.",
            "",
            "### Next step",
            "",
            "1. Have the doctor complete `doctor_validation_template.xlsx`",
            "   (the three highlighted columns only).",
            "2. Re-run:",
            "",
            "   ```bash",
            "   python tools/run_scoring_validation.py --doctor-file <completed file>",
            "   ```",
            "",
            "3. This report will then contain the full old-vs-new comparison.",
            "",
        ]
    else:
        lines += [
            "## 1. Overall result",
            "",
            "| Metric | Legacy | Canonical V2 | Delta | Better |",
            "| --- | --- | --- | --- | --- |",
        ]
        for entry in comparison:
            lines.append(
                f"| {entry['metric']} | {format_value(entry['legacy'])} | {format_value(entry['canonical_v2'])} "
                f"| {format_value(entry['delta'])} | {entry['better']} |"
            )
        lines += [
            "",
            "Delta is canonical minus legacy. For MAE, median error, large-bucket",
            "disagreement, false-safe and false-block counts a lower value is better;",
            "for every other metric a higher value is better.",
            "",
        ]

        improved = [e for e in comparison if e["better"] == "canonical_v2"]
        regressed = [e for e in comparison if e["better"] == "legacy"]
        lines += [
            "## 2. Where canonical V2 improved",
            "",
        ]
        lines += (
            [f"- **{e['metric']}**: {format_value(e['legacy'])} -> {format_value(e['canonical_v2'])}" for e in improved]
            if improved
            else ["- Nothing improved on the measured metrics."]
        )
        lines += ["", "## 3. Where canonical V2 regressed", ""]
        lines += (
            [f"- **{e['metric']}**: {format_value(e['legacy'])} -> {format_value(e['canonical_v2'])}" for e in regressed]
            if regressed
            else ["- Nothing regressed on the measured metrics."]
        )

        lines += ["", "## 4. Biggest failure modes", "", "| Failure mode | Canonical V2 | Legacy |", "| --- | --- | --- |"]
        for mode in modes:
            lines.append(f"| {mode['failure_mode']} | {mode['canonical_v2_cases']} | {mode['legacy_cases']} |")

        for title, key in [
            ("5. Product-type performance", "product_type"),
            ("6. Concern performance", "concern"),
            ("7. Safety performance (special condition)", "special_conditions"),
            ("8. Confidence performance", "confidence"),
        ]:
            rows_for_key = breakdown_files.get(key, [])
            lines += ["", f"## {title}", ""]
            if not rows_for_key:
                lines.append("- No data.")
                continue
            lines += [
                "| Group | Pairs | Legacy MAE | Canonical MAE | Better |",
                "| --- | --- | --- | --- | --- |",
            ]
            for row in rows_for_key:
                lines.append(
                    f"| {row[key]} | {format_value(row.get('canonical_pairs'))} "
                    f"| {format_value(row.get('legacy_mae'))} | {format_value(row.get('canonical_mae'))} "
                    f"| {row.get('better_mae')} |"
                )

        lines += ["", "## 9. Ranking performance", ""]
        if ranking:
            lines += [
                "| Profile | Rated | Legacy top-5 | Canonical top-5 | Legacy top-10 | Canonical top-10 |",
                "| --- | --- | --- | --- | --- | --- |",
            ]
            for row in ranking:
                lines.append(
                    f"| {row['profile_id']} | {row['rated_products']} "
                    f"| {format_value(row.get('legacy_top5_overlap_pct'))} "
                    f"| {format_value(row.get('canonical_top5_overlap_pct'))} "
                    f"| {format_value(row.get('legacy_top10_overlap_pct'))} "
                    f"| {format_value(row.get('canonical_top10_overlap_pct'))} |"
                )
        else:
            lines.append("- Not enough doctor-rated products per profile to rank.")

        lines += ["", "## 10. Top failures", "", f"The {len(failures)} largest canonical-vs-doctor errors are in `scoring_quality_failures.csv`.", ""]
        if failures:
            lines += [
                "| Product | Profile | Doctor | Canonical | Error | Suspected cause |",
                "| --- | --- | --- | --- | --- | --- |",
            ]
            for row in failures[:15]:
                lines.append(
                    f"| {row['product_name'][:40]} | {row['profile_id']} | {row['doctor_score']} "
                    f"| {row['canonical_v2_score']} | {row['absolute_error']} | {row['suspected_failure_category']} |"
                )

    lines += [
        "",
        "## 11. Limitations",
        "",
        "- **Divergent anchor configuration.** The shipped anchor layer uses the top 12",
        "  doctor anchors at similarity `>= 4.5`; the leave-one-out validation in",
        "  `tools/validate_v2_automated_logic.py` uses the top 5 at `>= 3.25`. These were",
        "  left divergent deliberately. This framework measures the **shipped**",
        "  configuration only, so its numbers are not comparable with the leave-one-out",
        "  report's numbers.",
        f"- **Legacy coverage.** Only {meta.get('legacy_matched_products')} of {meta.get('sample_size')} sampled",
        f"  products ({meta.get('legacy_coverage_pct')}%) exist in the legacy catalogue. The legacy",
        "  population was keyed by free-text product name, so the two catalogues are matched",
        "  by normalised name and the old-vs-new comparison is only valid on that subset.",
        "- **Ranking scope.** Ranking metrics are computed within the validation sample, not",
        "  the full catalogue. A ~150-product shelf is not the shelf a customer sees.",
        "- **Doctor coverage.** The doctor template is a prioritised subset of the full",
        "  product x profile grid, so per-group metrics rest on small counts. Treat any",
        "  breakdown group with few pairs as indicative only.",
        "- **Single reviewer.** With one doctor there is no inter-rater agreement estimate,",
        "  so doctor judgement is treated as ground truth without an error bar of its own.",
        "",
        "## 12. Recommendation for next step",
        "",
    ]

    if doctor_count == 0:
        lines += [
            "Collect the doctor review. No scoring change should be considered before",
            "there is evidence to justify it — the current systems are frozen and",
            "unmeasured, and any tuning now would be guesswork.",
            "",
        ]
    else:
        lines += [
            "Read the per-group breakdowns before drawing any conclusion from the overall",
            "table. An aggregate improvement can hide a safety regression: check the",
            "false-safe counts and the special-condition breakdown first, because those are",
            "the failures with clinical consequences rather than cosmetic ones.",
            "",
        ]

    lines += [
        "---",
        "",
        "*Generated by `tools/run_scoring_validation.py`. This framework measures only;",
        "it does not modify the scoring engine, its weights, or any production data.*",
        "",
    ]
    return "\n".join(lines)


def main() -> int:
    parser = argparse.ArgumentParser(description="Run scoring-quality validation metrics.")
    parser.add_argument("--doctor-file", type=Path, default=None, help="Completed doctor review file (xlsx or csv).")
    parser.add_argument("--failure-limit", type=int, default=50)
    args = parser.parse_args()

    rows = load_predictions()
    meta = json.loads(META_PATH.read_text(encoding="utf-8")) if META_PATH.exists() else {}

    if args.doctor_file:
        merged = merge_doctor_file(rows, args.doctor_file)
        print(f"Merged {merged} doctor answers from {args.doctor_file.name}")
        write_csv(PREDICTIONS_PATH, rows)

    doctor = [number(row["doctor_score"]) for row in rows]
    legacy_scores = [number(row["legacy_score"]) for row in rows]
    canonical_scores = [number(row["canonical_v2_score"]) for row in rows]

    doctor_count = sum(1 for value in doctor if value is not None)
    comparable = sum(
        1
        for d, l, c in zip(doctor, legacy_scores, canonical_scores)
        if d is not None and l is not None and c is not None
    )

    print(f"Predictions: {len(rows)} pairs")
    print(f"Doctor scores present: {doctor_count}")
    print(f"Comparable across both systems + doctor: {comparable}")

    # Fair comparison: restrict both suites to the pairs where BOTH systems have
    # a score, so legacy is not judged on a different set of products.
    paired_rows = [
        row
        for row in rows
        if number(row["doctor_score"]) is not None
        and number(row["legacy_score"]) is not None
        and number(row["canonical_v2_score"]) is not None
    ]
    legacy_suite = sv.metric_suite(
        [number(row["legacy_score"]) for row in paired_rows],
        [number(row["doctor_score"]) for row in paired_rows],
    )
    canonical_suite = sv.metric_suite(
        [number(row["canonical_v2_score"]) for row in paired_rows],
        [number(row["doctor_score"]) for row in paired_rows],
    )
    # Canonical measured over everything the doctor rated, for context.
    canonical_all = sv.metric_suite(canonical_scores, doctor)
    comparison = sv.compare_suites(legacy_suite, canonical_suite)

    breakdown_files: dict[str, list[dict[str, Any]]] = {}
    for key, filename in BREAKDOWNS:
        table = sv.breakdown(rows, key)
        breakdown_files[key] = table
        write_csv(CANONICAL_OUTPUT_DIR / filename, table)

    ranking = ranking_report(rows)
    write_csv(RANKING_PATH, ranking)

    failures = failure_analysis(rows, args.failure_limit)
    write_csv(FAILURES_PATH, failures)

    modes = known_failure_modes(rows)

    METRICS_PATH.write_text(
        json.dumps(
            {
                "scoring_validation_version": sv.SCORING_VALIDATION_VERSION,
                "generated_at": datetime.now(timezone.utc).isoformat(),
                "doctor_scores_present": doctor_count,
                "total_pairs": len(rows),
                "comparable_pairs": comparable,
                "legacy": legacy_suite,
                "canonical_v2": canonical_suite,
                "canonical_v2_all_rated_pairs": canonical_all,
                "comparison": comparison,
                "known_failure_modes": modes,
                "ranking": ranking,
                "methodology": meta.get("metric_methodology"),
                "limitations": meta.get("known_limitations"),
            },
            indent=2,
        ),
        encoding="utf-8",
    )

    report = build_report(
        meta, comparison, legacy_suite, canonical_suite, breakdown_files, ranking, failures, modes,
        doctor_count, len(rows), comparable,
    )
    REPORT_PATH.write_text(report, encoding="utf-8")

    print(f"\nReport  -> {REPORT_PATH.name}")
    print(f"Metrics -> {METRICS_PATH.name}")
    for _key, filename in BREAKDOWNS:
        print(f"Breakdown -> {filename}")
    print(f"Ranking -> {RANKING_PATH.name}")
    print(f"Failures -> {FAILURES_PATH.name}")

    if doctor_count == 0:
        print(
            "\nNo doctor scores present. Metrics are undefined and the report says so.\n"
            "No conclusion about either system can be drawn yet."
        )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
